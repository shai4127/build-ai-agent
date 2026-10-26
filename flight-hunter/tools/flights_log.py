"""Maintain flights.xlsx: one row per price check.

Commands:
    python tools/flights_log.py init
    python tools/flights_log.py append --json "<json>"      (or pipe JSON on stdin)
    python tools/flights_log.py status

append JSON fields (all optional except notes when nothing was found):
    site, airline, out_flight, out_times, ret_flight, ret_times,
    base_price, bag_fee, total_price, link, notes

If total_price is missing/null the row is recorded as "לא נמצא".
The script stamps the check date and time (Israel time) itself, computes the change
from the previous priced row and reports whether this is a new lowest price.

If flights.xlsx is open in Excel (locked), the row is queued in pending_rows.jsonl
and written automatically on the next successful append.
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
CONFIG = json.loads((ROOT / "tools" / "config.json").read_text(encoding="utf-8"))
XLSX = ROOT / CONFIG["excel_file"]
PENDING = ROOT / "pending_rows.jsonl"
SHEET = "כל הבדיקות"
TZ = ZoneInfo(CONFIG["timezone_display"])

HEADERS = [
    "תאריך בדיקה",
    "שעת בדיקה (שעון ישראל)",
    "אתר",
    "חברת תעופה",
    "מס' טיסה הלוך",
    "המראה/נחיתה הלוך",
    "מס' טיסה חזור",
    "המראה/נחיתה חזור",
    "מחיר בסיס (₪)",
    "תוספת מזוודה (₪)",
    "מחיר כולל לשניים (₪)",
    "שינוי מהבדיקה הקודמת (₪)",
    "קישור",
    "הערות",
]
COL_TOTAL = HEADERS.index("מחיר כולל לשניים (₪)") + 1
WIDTHS = [13, 12, 16, 18, 14, 22, 14, 22, 14, 14, 18, 18, 40, 45]


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return wb


def _load() -> Workbook:
    if XLSX.exists():
        return load_workbook(XLSX)
    return _new_workbook()


def _num(value):
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace(",", "").replace("₪", "").strip()), 2)
    except ValueError:
        return None


def _priced_totals(ws) -> list[float]:
    totals = []
    for row in ws.iter_rows(min_row=2, min_col=COL_TOTAL, max_col=COL_TOTAL, values_only=True):
        value = _num(row[0])
        if value is not None:
            totals.append(value)
    return totals


def _write_row(ws, data: dict, stamp: datetime) -> dict:
    totals = _priced_totals(ws)
    previous_total = totals[-1] if totals else None
    previous_best = min(totals) if totals else None

    total = _num(data.get("total_price"))
    notes = (data.get("notes") or "").strip()
    if total is None and "לא נמצא" not in notes:
        notes = ("לא נמצא. " + notes).strip()

    change = None
    if total is not None and previous_total is not None:
        change = round(total - previous_total, 2)

    ws.append([
        stamp.strftime("%d/%m/%Y"),
        stamp.strftime("%H:%M"),
        data.get("site") or "",
        data.get("airline") or "",
        data.get("out_flight") or "",
        data.get("out_times") or "",
        data.get("ret_flight") or "",
        data.get("ret_times") or "",
        _num(data.get("base_price")),
        _num(data.get("bag_fee")),
        total,
        change,
        data.get("link") or "",
        notes,
    ])
    row_idx = ws.max_row
    link = data.get("link")
    if link:
        cell = ws.cell(row=row_idx, column=HEADERS.index("קישור") + 1)
        cell.hyperlink = link
        cell.font = Font(color="0563C1", underline="single")
    for col in (9, 10, 11, 12):
        ws.cell(row=row_idx, column=col).number_format = "#,##0"

    return {
        "row": row_idx,
        "checked_at": stamp.strftime("%d/%m/%Y %H:%M"),
        "total_price": total,
        "previous_total": previous_total,
        "change": change,
        "previous_best": previous_best,
        "first_price": total is not None and previous_best is None,
        "new_low": total is not None and previous_best is not None and total < previous_best,
    }


def cmd_init() -> dict:
    if XLSX.exists():
        return {"created": False, "file": str(XLSX)}
    _new_workbook().save(XLSX)
    return {"created": True, "file": str(XLSX)}


def cmd_append(data: dict) -> dict:
    stamp = datetime.now(TZ)
    queued = []
    if PENDING.exists():
        queued = [json.loads(line) for line in PENDING.read_text(encoding="utf-8").splitlines() if line.strip()]

    wb = _load()
    ws = wb[SHEET]
    for item in queued:
        _write_row(ws, item["data"], datetime.fromisoformat(item["stamp"]))
    result = _write_row(ws, data, stamp)
    try:
        wb.save(XLSX)
    except PermissionError:
        with PENDING.open("a", encoding="utf-8") as f:
            f.write(json.dumps({"stamp": stamp.isoformat(), "data": data}, ensure_ascii=False) + "\n")
        result["saved"] = False
        result["warning"] = "flights.xlsx is open/locked; row queued in pending_rows.jsonl"
        return result
    if PENDING.exists():
        PENDING.unlink()
    result["saved"] = True
    result["flushed_pending_rows"] = len(queued)
    return result


def cmd_status() -> dict:
    if not XLSX.exists():
        return {"exists": False}
    ws = load_workbook(XLSX, read_only=True)[SHEET]
    totals = _priced_totals(ws)
    return {
        "exists": True,
        "checks": ws.max_row - 1,
        "priced_checks": len(totals),
        "last_total": totals[-1] if totals else None,
        "best_total": min(totals) if totals else None,
        "pending_rows": PENDING.exists(),
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init")
    sub.add_parser("status")
    ap = sub.add_parser("append")
    ap.add_argument("--json", help="row data as JSON; read from stdin when omitted")
    args = parser.parse_args()

    if args.cmd == "init":
        out = cmd_init()
    elif args.cmd == "status":
        out = cmd_status()
    else:
        raw = args.json if args.json else sys.stdin.read()
        out = cmd_append(json.loads(raw))
    print(json.dumps(out, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
